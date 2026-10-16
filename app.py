import pandas as pd
import streamlit as st
import plotly.express as px
import datetime
import os
import openpyxl

# --- PAGE CONFIGURATION ---
st.set_page_config(
    page_title="Indian Railways — TADK Executive BI Suite",
    page_icon="🚂",
    layout="wide",
    initial_sidebar_state="expanded"
)

# --- ADVANCED STYLING ---
st.markdown("""
    <style>
    .main { background-color: #0b0f19; }
    .stMetric {
        background: linear-gradient(135deg, rgba(30, 41, 59, 0.7) 0%, rgba(15, 23, 42, 0.9) 100%);
        border: 1px solid rgba(56, 189, 248, 0.2);
        padding: 20px;
        border-radius: 14px;
    }
    h1, h2, h3 { font-family: 'Inter', sans-serif; color: #f1f5f9; }
    </style>
""", unsafe_allow_html=True)

# --- SIMPLIFIED ROLE-BASED AUTHENTICATION SYSTEM ---
USERS = {
    "admin": {"password": "railway123", "role": "Administrator", "zone": "ALL"},
    "dnr": {"password": "123", "role": "Zonal Officer", "zone": "DNR"},
    "ddu": {"password": "123", "role": "Zonal Officer", "zone": "DDU"},
    "dhn": {"password": "123", "role": "Zonal Officer", "zone": "DHN"},
    "see": {"password": "123", "role": "Zonal Officer", "zone": "SEE"},
    "spj": {"password": "123", "role": "Zonal Officer", "zone": "SPJ"}
}

def check_authentication():
    if "authenticated" not in st.session_state:
        st.session_state["authenticated"] = False
        st.session_state["username"] = ""
        st.session_state["role"] = ""

    if not st.session_state["authenticated"]:
        st.markdown("<br><br>", unsafe_allow_html=True)
        col1, col2, col3 = st.columns([1, 2, 1])
        with col2:
            st.markdown("## 🔒 Indian Railways Secure Login")
            st.markdown("Enter credentials to access the TADK Executive BI Suite.")
            username_input = st.text_input("Username").strip().lower()
            password_input = st.text_input("Password", type="password")

            if st.button("Login", type="primary"):
                if username_input in USERS and USERS[username_input]["password"] == password_input:
                    st.session_state["authenticated"] = True
                    st.session_state["username"] = username_input
                    st.session_state["role"] = USERS[username_input]["role"]
                    st.session_state["assigned_zone"] = USERS[username_input]["zone"]
                    st.rerun()
                else:
                    st.error("❌ Invalid Username or Password")
        return False
    return True

if not check_authentication():
    st.stop()

# --- SIDEBAR USER SESSION INFO ---
st.sidebar.markdown(f"👤 **Logged in as:** `{st.session_state['username'].upper()}`")
st.sidebar.markdown(f"🛡️ **Role:** `{st.session_state['role']}`")
if st.sidebar.button("Logout"):
    st.session_state["authenticated"] = False
    st.rerun()

st.sidebar.markdown("---")

# --- STANDARDIZED TEMPLATE SETUP ---
TEMPLATE_COLUMNS = [
    'Sl. No',
    'Department',
    'Total Posts',
    'Attached Post',
    'Post Type',
    'Currency Upto',
    'Incumbent Officer',
    'Engagement Type'
]

DEFAULT_DEPARTMENTS = [
    'Engineering', 'Electrical', 'Mechanical', 'Signal & Telecom',
    'Operating', 'Commercial', 'Account', 'Personnel', 'Medical', 'General Admin', 'Store'
]

DEFAULT_POST_TYPES = ['Revenue (Permanent/Temp)', 'Work Charge', 'Other']
DEFAULT_ENGAGEMENT_TYPES = ['Railway Regular', 'Contractual', 'Vacant / Unspecified', 'Other / Processing']

excel_file = "New TADK Position as on 01.10.2026.xlsx"

# --- SELF-HEALING EXCEL INITIALIZER ---
def ensure_excel_file():
    if not os.path.exists(excel_file):
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "DNR"
        ws.append(TEMPLATE_COLUMNS)
        wb.save(excel_file)

ensure_excel_file()

@st.cache_data
def load_and_clean_all_sheets(file_path):
    try:
        xls = pd.ExcelFile(file_path, engine='openpyxl')
    except Exception:
        ensure_excel_file()
        xls = pd.ExcelFile(file_path, engine='openpyxl')

    all_data = []

    for sheet in xls.sheet_names:
        if sheet.lower() == 'sheet':
            continue
        try:
            df = pd.read_excel(file_path, sheet_name=sheet, engine='openpyxl')
        except Exception:
            continue

        # Ignore empty sheets without data rows
        if df.shape[0] >= 1 and len(df.dropna(how='all')) > 0:
            raw_headers = df.columns.values
            headers = []
            for i, h in enumerate(raw_headers):
                if pd.isna(h):
                    headers.append(f"Extra_{i}")
                else:
                    headers.append(str(h).strip().replace('\n', ' '))

            sub_df = df.copy()
            sub_df.columns = headers[:sub_df.shape[1]]
            sub_df['Sheet_Name'] = sheet

            parts = sheet.split('_')
            if len(parts) >= 2:
                sub_df['Zone_Division'] = parts[0]
                sub_df['Reporting_Month'] = "_".join(parts[1:])
            else:
                sub_df['Zone_Division'] = sheet
                sub_df['Reporting_Month'] = "Oct_2026"

            if len(sub_df.columns) > 1:
                dept_col = sub_df.columns[1]
                sub_df[dept_col] = sub_df[dept_col].ffill()

            all_data.append(sub_df)

    if all_data:
        standardized_dfs = []
        for d in all_data:
            cols = list(d.columns)
            temp_df = pd.DataFrame()
            temp_df['Sl. No'] = d.iloc[:, 0] if len(cols) > 0 else 1
            temp_df['Department'] = d.iloc[:, 1] if len(cols) > 1 else 'Engineering'
            temp_df['Total Posts'] = d.iloc[:, 2] if len(cols) > 2 else 1
            temp_df['Attached Post'] = d.iloc[:, 3] if len(cols) > 3 else 'Unspecified'
            temp_df['Post Type'] = d.iloc[:, 4] if len(cols) > 4 else 'Revenue (Permanent/Temp)'
            temp_df['Currency Upto'] = d.iloc[:, 5] if len(cols) > 5 else 'Permanent'
            temp_df['Incumbent Officer'] = d.iloc[:, 6] if len(cols) > 6 else 'Vacant'
            temp_df['Engagement Type'] = d.iloc[:, 7] if len(cols) > 7 else 'Vacant / Unspecified'
            temp_df['Zone_Division'] = d['Zone_Division']
            temp_df['Reporting_Month'] = d['Reporting_Month']
            temp_df['Sheet_Name'] = d['Sheet_Name']
            standardized_dfs.append(temp_df)

        combined = pd.concat(standardized_dfs, ignore_index=True)
        combined = combined[~combined['Department'].astype(str).str.lower().str.contains('total', na=False)]
        combined = combined[~combined['Sl. No'].astype(str).str.lower().str.contains('total', na=False)]

        combined['Department'] = combined['Department'].fillna('Engineering').astype(str).str.title()
        combined['Post Type'] = combined['Post Type'].fillna('Revenue (Permanent/Temp)').astype(str).str.strip()
        combined['Engagement Type'] = combined['Engagement Type'].fillna('Vacant / Unspecified').astype(str).str.strip()

        def clean_engagement(val):
            s = str(val).strip().lower()
            if any(k in s for k in ['vacant', 'vaccant', 'nil', '...', '---', 'nan', 'unspecified']):
                return 'Vacant / Unspecified'
            elif 'contract' in s or 'attendant' in s:
                return 'Contractual'
            elif 'railway' in s or 'tadk' in s:
                return 'Railway Regular'
            else:
                return 'Other / Processing'

        combined['Status'] = combined['Engagement Type'].apply(clean_engagement)

        def clean_post_type(val):
            s = str(val).strip().lower()
            if 'work' in s:
                return 'Work Charge'
            elif 'revenue' in s:
                return 'Revenue (Permanent/Temp)'
            else:
                return 'Other'

        combined['Structure'] = combined['Post Type'].apply(clean_post_type)
        return combined
    return pd.DataFrame()

df = load_and_clean_all_sheets(excel_file)

if df.empty:
    st.warning("No data records found in the sheets yet. Please use the Zonal Editor tab to add records.")
else:
    # --- SIDEBAR FILTERS ---
    st.sidebar.header("🎛️ Executive Filters & Control")

    all_months = sorted(df['Reporting_Month'].unique().tolist())
    selected_months = st.sidebar.multiselect("🗓️ Select Reporting Month(s):", options=all_months, default=all_months)

    filtered_by_month = df[df['Reporting_Month'].isin(selected_months)]

    if st.session_state['role'] == 'Zonal Officer':
        all_zones = [st.session_state['assigned_zone']]
        selected_zones = st.sidebar.multiselect("🏢 Select Zone / Division(s):", options=all_zones, default=all_zones, disabled=True)
    else:
        all_zones = sorted(filtered_by_month['Zone_Division'].unique().tolist())
        selected_zones = st.sidebar.multiselect("🏢 Select Zone / Division(s):", options=all_zones, default=all_zones)

    filtered_by_zone = filtered_by_month[filtered_by_month['Zone_Division'].isin(selected_zones)]

    all_depts = sorted(filtered_by_zone['Department'].unique().tolist())
    selected_depts = st.sidebar.multiselect("📁 Select Department(s):", options=all_depts, default=all_depts)

    all_status = sorted(filtered_by_zone['Status'].unique().tolist())
    selected_status = st.sidebar.multiselect("📌 Select Engagement Status:", options=all_status, default=all_status)

    search_query = st.sidebar.text_input("🔍 Search Officer, Post, or Dept:", "").strip().lower()

    filtered_df = filtered_by_zone[
        filtered_by_zone['Department'].isin(selected_depts) &
        filtered_by_zone['Status'].isin(selected_status)
    ]

    if search_query:
        mask = filtered_df.astype(str).apply(lambda x: x.str.lower().str.contains(search_query)).any(axis=1)
        filtered_df = filtered_df[mask]

# --- MAIN HEADER ---
st.title("🚂 Indian Railways — TADK Executive BI Suite")
st.markdown("### Real-time Zonal Intelligence, Monthly Reporting & Live Operations")
st.markdown("---")

# --- TABS ---
tab1, tab2, tab3 = st.tabs(["📊 Executive Analytics & Metrics", "📋 Master Records Explorer", "✏️ Monthly Data Entry & Zonal Editor"])

with tab1:
    if df.empty:
        st.info("No data available to visualize. Please add records via Tab 3.")
    else:
        m1, m2, m3, m4 = st.columns(4)
        total_pos = len(filtered_df)
        contractual_pos = len(filtered_df[filtered_df['Status'] == 'Contractual'])
        vacant_pos = len(filtered_df[filtered_df['Status'] == 'Vacant / Unspecified'])
        railway_pos = len(filtered_df[filtered_df['Status'] == 'Railway Regular'])

        m1.metric("Filtered Total Posts", total_pos)
        m2.metric("Contractual Share", contractual_pos, delta=f"{round(contractual_pos/total_pos*100, 1) if total_pos else 0}%")
        m3.metric("Vacant / Unspecified", vacant_pos, delta_color="inverse")
        m4.metric("Railway Regular", railway_pos)

        st.markdown("---")

        c1, c2 = st.columns(2)
        with c1:
            st.subheader("📊 Zone/Division Distribution")
            if not filtered_df.empty:
                zone_counts = filtered_df['Zone_Division'].value_counts().reset_index()
                zone_counts.columns = ['Zone_Division', 'Count']
                fig_zone = px.bar(zone_counts, x='Zone_Division', y='Count', color='Zone_Division', text='Count', template='plotly_dark')
                st.plotly_chart(fig_zone, use_container_width=True)
            else:
                st.info("No records match the current filter.")

        with c2:
            st.subheader("🍩 Engagement Status Distribution")
            if not filtered_df.empty:
                status_counts = filtered_df['Status'].value_counts().reset_index()
                status_counts.columns = ['Status', 'Count']
                fig_status = px.pie(status_counts, names='Status', values='Count', hole=0.6, template='plotly_dark')
                st.plotly_chart(fig_status, use_container_width=True)
            else:
                st.info("No records match the current filter.")

with tab2:
    st.subheader("📋 Filtered Active Records Explorer")
    if not df.empty:
        csv_data = filtered_df.to_csv(index=False).encode('utf-8')
        st.download_button(label="📥 Export Filtered CSV", data=csv_data, file_name="filtered_tadk_executive_report.csv", mime="text/csv")
        st.dataframe(filtered_df[['Reporting_Month', 'Zone_Division', 'Department', 'Total Posts', 'Attached Post', 'Post Type', 'Currency Upto', 'Incumbent Officer', 'Engagement Type', 'Status']], use_container_width=True, height=500)
    else:
        st.info("No records found.")

with tab3:
    st.subheader("✏️ Monthly Zonal Data Entry & Management Suite")

    wb_check = openpyxl.load_workbook(excel_file)
    sheet_list = wb_check.sheetnames

    base_zones = sorted(list(set([s.split('_')[0] for s in sheet_list if '_' in s])))
    if not base_zones:
        base_zones = ['DNR', 'DDU', 'DHN', 'SEE', 'SPJ', 'HQ']

    # ADMIN ONLY FEATURE: Batch Create Brand New Blank Month Suite for ALL Zones
    if st.session_state['role'] == 'Administrator':
        with st.expander("🛠️ [Admin Only] Initialize New Month Suite for All Zones", expanded=False):
            st.markdown("Entering a new month here will automatically create **completely blank, pristine sheets** for all zones at once with standard pre-headers.")
            new_month_name = st.text_input("Enter New Reporting Month (e.g., Nov_2026):").strip()
            if st.button("🚀 Generate Blank Month Suite for All Zones", type="primary"):
                if new_month_name:
                    try:
                        wb = openpyxl.load_workbook(excel_file)
                        for zone in base_zones:
                            new_sheet_name = f"{zone}_{new_month_name}"
                            if new_sheet_name not in wb.sheetnames:
                                ws = wb.create_sheet(title=new_sheet_name)
                                ws.append(TEMPLATE_COLUMNS)
                        wb.save(excel_file)

                        st.cache_data.clear()
                        st.success(f"✅ Successfully created blank month suite **{new_month_name}** for all zones ({', '.join(base_zones)})!")
                        st.rerun()
                    except Exception as e:
                        st.error(f"❌ Error creating month suite: {e}")
                else:
                    st.warning("Please enter a valid month name.")
        st.markdown("---")

    if st.session_state['role'] == 'Zonal Officer':
        allowed_sheets = [s for s in sheet_list if s.startswith(st.session_state['assigned_zone'])]
        selected_sheet = st.selectbox("Select Your Assigned Division & Month Sheet:", options=allowed_sheets)
    else:
        selected_sheet = st.selectbox("Select Zone / Division & Month Sheet to Edit:", options=sheet_list)

    df_sheet = pd.read_excel(excel_file, sheet_name=selected_sheet, engine='openpyxl')

    for col in TEMPLATE_COLUMNS:
        if col not in df_sheet.columns:
            df_sheet[col] = ''

    st.markdown(f"### Editing Sheet: `{selected_sheet}`")
    st.info("💡 **Clean Slate Active:** New sheets start completely blank. Use the table below to add entries, assign officers, and track positions.")

    edited_df = st.data_editor(
        df_sheet[TEMPLATE_COLUMNS],
        num_rows="dynamic",
        use_container_width=True,
        height=450,
        column_config={
            "Department": st.column_config.SelectboxColumn("Department", options=DEFAULT_DEPARTMENTS, required=True),
            "Post Type": st.column_config.SelectboxColumn("Post Type", options=DEFAULT_POST_TYPES, required=True),
            "Engagement Type": st.column_config.SelectboxColumn("Engagement Type", options=DEFAULT_ENGAGEMENT_TYPES, required=True),
            "Total Posts": st.column_config.NumberColumn("Total Posts", min_value=1, max_value=10, step=1, required=True)
        }
    )

    if st.button("💾 Save & Commit Updates to Master File", type="primary"):
        try:
            wb = openpyxl.load_workbook(excel_file)
            if selected_sheet in wb.sheetnames:
                ws = wb[selected_sheet]
                ws.delete_rows(1, ws.max_row)
                ws.append(TEMPLATE_COLUMNS)
                cleaned_edited_df = edited_df.dropna(how='all')
                for _, row in cleaned_edited_df.iterrows():
                    ws.append(list(row))
                wb.save(excel_file)

            st.cache_data.clear()
            timestamp = datetime.datetime.now().strftime("%Y-%m-%d %H:%M:%S")
            st.success(f"✅ Successfully saved **{selected_sheet}** on {timestamp}! All analytics have been instantly refreshed.")
        except Exception as e:
            st.error(f"❌ Error saving sheet: {e}")
